#!/usr/bin/env python3
"""Export selected or weekly trips from Excel into the supplied settlement PDF."""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from pypdf import PdfReader, PdfWriter
from pypdf.generic import NameObject, TextStringObject


WEEKLY_ROW_START = 11
WEEKLY_ROW_END = 30


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value: Any, default: float = 0.0) -> float:
    if value in (None, ""):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("$", "").replace(",", "").strip()
    return float(text) if text else default


def whole(value: Any) -> int:
    return int(round(number(value)))


def money(value: float, blank_zero: bool = True) -> str:
    if blank_zero and abs(value) < 0.005:
        return ""
    return f"{value:,.2f}"


def decimal(value: Any) -> str:
    if value in (None, ""):
        return ""
    n = number(value)
    return f"{n:.2f}".rstrip("0").rstrip(".")


def as_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except Exception:
            return None
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def pdf_date(value: Any) -> str:
    d = as_date(value)
    return d.strftime("%m-%d-%Y") if d else ""


def short_date(value: Any) -> str:
    d = as_date(value)
    return d.strftime("%m/%d/%y") if d else ""


def safe_name(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return value.strip("_") or "Driver"


def same_trip(a: Any, b: Any) -> bool:
    return clean(a).lstrip("0") == clean(b).lstrip("0") and clean(a) != "" and clean(b) != ""


@dataclass
class Trip:
    trip_number: str
    driver: str
    co_driver: str
    pickup_date: Any
    delivery_date: Any
    pro_bol: str
    truck: str
    trailer: str
    origin: str
    destination: str
    begin_odo: int
    end_odo: int
    actual_miles: int
    paid_miles: int
    rate: float
    mileage_pay: float
    extra_stop_pay: float
    layover_pay: float
    other_pay: float
    gross_pay: float
    reimbursements: float
    expenses: float
    advances: float
    tax_reserve_rate: float
    routes: list[dict[str, Any]]
    states: list[dict[str, Any]]
    fuel: list[dict[str, Any]]
    expense_rows: list[dict[str, Any]]
    advance_rows: list[dict[str, Any]]
    warnings: list[str]

    @property
    def state_miles(self) -> int:
        return sum(whole(row["miles"]) for row in self.states)

    @property
    def fuel_total(self) -> float:
        return sum(number(row["total"]) for row in self.fuel)

    @property
    def settlement_pay(self) -> float:
        return self.gross_pay + self.reimbursements - self.advances


class WorkbookData:
    def __init__(self, path: Path):
        self.path = path
        self.book = load_workbook(path, data_only=False, read_only=False)
        self.weekly = self.book["Weekly Pay"]

    def selected_trip(self) -> str:
        return clean(self.weekly["N4"].value)

    def all_trip_numbers(self) -> list[str]:
        trips: list[str] = []
        for row in range(WEEKLY_ROW_START, WEEKLY_ROW_END + 1):
            value = clean(self.weekly.cell(row=row, column=3).value)
            if value and value not in trips:
                trips.append(value)
        return trips

    def _rows(self, sheet_name: str, start: int, end: int, trip_col: int, trip_number: str):
        sheet = self.book[sheet_name]
        return [row for row in range(start, end + 1) if same_trip(sheet.cell(row=row, column=trip_col).value, trip_number)]

    def trip(self, trip_number: str) -> Trip:
        row = next((r for r in range(WEEKLY_ROW_START, WEEKLY_ROW_END + 1) if same_trip(self.weekly.cell(r, 3).value, trip_number)), None)
        if row is None:
            raise ValueError(f"Trip {trip_number} was not found on Weekly Pay.")

        begin_raw = self.weekly.cell(row, 9).value
        end_raw = self.weekly.cell(row, 10).value
        paid_raw = self.weekly.cell(row, 12).value
        rate_raw = self.weekly.cell(row, 13).value
        missing = []
        if begin_raw in (None, ""): missing.append("beginning odometer")
        if end_raw in (None, ""): missing.append("ending odometer")
        if paid_raw in (None, ""): missing.append("paid miles")
        if rate_raw in (None, ""): missing.append("mileage rate")
        if missing:
            raise ValueError(f"Trip {trip_number} is missing: {', '.join(missing)}.")

        begin_odo = whole(begin_raw)
        end_odo = whole(end_raw)
        if end_odo < begin_odo:
            raise ValueError(f"Trip {trip_number}: ending odometer cannot be lower than beginning odometer.")
        actual_miles = end_odo - begin_odo
        paid_miles = whole(paid_raw)
        rate = number(rate_raw)
        if rate <= 0:
            raise ValueError(f"Trip {trip_number}: mileage rate must be greater than zero.")

        route_rows = self._rows("Routes", 7, 26, 1, trip_number)
        routes = [
            {
                "order": whole(self.book["Routes"].cell(r, 2).value or 999),
                "type": clean(self.book["Routes"].cell(r, 3).value),
                "from": clean(self.book["Routes"].cell(r, 4).value),
                "to": clean(self.book["Routes"].cell(r, 5).value),
                "notes": clean(self.book["Routes"].cell(r, 6).value),
            }
            for r in route_rows
            if clean(self.book["Routes"].cell(r, 4).value) or clean(self.book["Routes"].cell(r, 5).value)
        ]
        routes.sort(key=lambda item: item["order"])

        state_rows = self._rows("State Miles", 7, 56, 1, trip_number)
        states = [
            {
                "state": clean(self.book["State Miles"].cell(r, 2).value).upper(),
                "miles": whole(self.book["State Miles"].cell(r, 3).value),
                "highways": clean(self.book["State Miles"].cell(r, 4).value),
            }
            for r in state_rows
            if clean(self.book["State Miles"].cell(r, 2).value)
        ]
        if not states:
            raise ValueError(f"Trip {trip_number}: state miles are required before export.")

        fuel_rows = self._rows("Fuel & DEF", 7, 36, 2, trip_number)
        fuel = [
            {
                "date": self.book["Fuel & DEF"].cell(r, 1).value,
                "state": clean(self.book["Fuel & DEF"].cell(r, 3).value).upper(),
                "invoice": clean(self.book["Fuel & DEF"].cell(r, 4).value),
                "seller": clean(self.book["Fuel & DEF"].cell(r, 5).value),
                "truck": self.book["Fuel & DEF"].cell(r, 6).value,
                "reefer": self.book["Fuel & DEF"].cell(r, 7).value,
                "def": self.book["Fuel & DEF"].cell(r, 8).value,
                "total": self.book["Fuel & DEF"].cell(r, 9).value,
            }
            for r in fuel_rows
        ]

        expense_rows_index = self._rows("Expenses", 7, 56, 2, trip_number)
        expense_rows = [
            {
                "date": self.book["Expenses"].cell(r, 1).value,
                "category": clean(self.book["Expenses"].cell(r, 3).value),
                "description": clean(self.book["Expenses"].cell(r, 4).value),
                "state": clean(self.book["Expenses"].cell(r, 5).value).upper(),
                "receipt": clean(self.book["Expenses"].cell(r, 6).value),
                "amount": self.book["Expenses"].cell(r, 7).value,
                "reimbursement": self.book["Expenses"].cell(r, 8).value,
            }
            for r in expense_rows_index
        ]

        advance_rows_index = self._rows("Advances", 7, 26, 2, trip_number)
        advance_rows = [
            {
                "date": self.book["Advances"].cell(r, 1).value,
                "amount": self.book["Advances"].cell(r, 3).value,
                "type": clean(self.book["Advances"].cell(r, 4).value),
                "notes": clean(self.book["Advances"].cell(r, 5).value),
            }
            for r in advance_rows_index
            if self.book["Advances"].cell(r, 3).value not in (None, "") or clean(self.book["Advances"].cell(r, 4).value)
        ]

        expenses = sum(number(item["amount"]) for item in expense_rows)
        reimbursements = sum(number(item["reimbursement"]) for item in expense_rows)
        advances = sum(number(item["amount"]) for item in advance_rows)
        extra = number(self.weekly.cell(row, 15).value)
        layover = number(self.weekly.cell(row, 16).value)
        other = number(self.weekly.cell(row, 17).value)
        mileage_pay = paid_miles * rate
        gross = mileage_pay + extra + layover + other

        warnings: list[str] = []
        state_total = sum(item["miles"] for item in states)
        if state_total != actual_miles:
            warnings.append(f"State miles ({state_total:,}) differ from actual odometer miles ({actual_miles:,}) by {state_total - actual_miles:+,}.")
        if paid_miles != actual_miles:
            warnings.append(f"Paid miles ({paid_miles:,}) differ from actual odometer miles ({actual_miles:,}) by {paid_miles - actual_miles:+,}.")
        for item in expense_rows:
            if item["description"] and item["amount"] in (None, ""):
                warnings.append(f"Expense amount missing: {item['description']}.")
            if number(item["amount"]) > 0 and not item["receipt"]:
                warnings.append(f"Receipt/invoice missing for expense: {item['description'] or item['category']}.")
        for item in fuel:
            if number(item["total"]) > 0 and not item["invoice"]:
                warnings.append(f"Invoice number missing for fuel purchase from {item['seller'] or 'unknown seller'}.")

        return Trip(
            trip_number=clean(self.weekly.cell(row, 3).value),
            driver=clean(self.weekly["B4"].value) or "Dustin Douglas",
            co_driver="N/A",
            pickup_date=self.weekly.cell(row, 1).value,
            delivery_date=self.weekly.cell(row, 2).value,
            pro_bol=clean(self.weekly.cell(row, 4).value),
            truck=clean(self.weekly.cell(row, 5).value),
            trailer=clean(self.weekly.cell(row, 6).value),
            origin=clean(self.weekly.cell(row, 7).value),
            destination=clean(self.weekly.cell(row, 8).value),
            begin_odo=begin_odo,
            end_odo=end_odo,
            actual_miles=actual_miles,
            paid_miles=paid_miles,
            rate=rate,
            mileage_pay=mileage_pay,
            extra_stop_pay=extra,
            layover_pay=layover,
            other_pay=other,
            gross_pay=gross,
            reimbursements=reimbursements,
            expenses=expenses,
            advances=advances,
            tax_reserve_rate=number(self.weekly["K4"].value),
            routes=routes,
            states=states,
            fuel=fuel,
            expense_rows=expense_rows,
            advance_rows=advance_rows,
            warnings=warnings,
        )


def category_amount(trip: Trip, *names: str) -> float:
    wanted = {name.lower() for name in names}
    return sum(number(item["amount"]) for item in trip.expense_rows if item["category"].strip().lower() in wanted)


def build_field_values(trip: Trip, field_names: set[str]) -> dict[str, str]:
    values = {name: "" for name in field_names}
    values.update({
        "driver_name": trip.driver,
        "co_driver_name": trip.co_driver,
        "truck_number": trip.truck,
        "trailer_number": trip.trailer,
        "trip_number": trip.trip_number,
        "beginning_odometer": str(trip.begin_odo),
        "ending_odometer": str(trip.end_odo),
        "start_date": pdf_date(trip.pickup_date),
        "end_date": pdf_date(trip.delivery_date),
    })

    routes = trip.routes or ([{"from": trip.origin, "to": trip.destination}] if trip.origin or trip.destination else [])
    for index, item in enumerate(routes[:10], 1):
        values[f"route_{index}_from"] = item.get("from", "")
        values[f"route_{index}_to"] = item.get("to", "")

    for index, item in enumerate(trip.states[:40], 1):
        side = "left" if index <= 20 else "right"
        slot = index if index <= 20 else index - 20
        values[f"{side}_state_{slot}"] = item["state"]
        values[f"{side}_miles_{slot}"] = str(item["miles"])
        values[f"{side}_highways_{slot}"] = item["highways"]
    values["left_state_miles_total"] = f"{trip.state_miles:,}"
    values["right_state_miles_total"] = ""

    for index, item in enumerate(trip.fuel[:20], 1):
        values[f"fuel_{index}_date"] = short_date(item["date"])
        values[f"fuel_{index}_state"] = item["state"]
        values[f"fuel_{index}_invoice"] = item["invoice"]
        values[f"fuel_{index}_seller"] = item["seller"]
        values[f"fuel_{index}_truck_gallons"] = decimal(item["truck"])
        values[f"fuel_{index}_reefer_gallons"] = decimal(item["reefer"])
        values[f"fuel_{index}_def"] = decimal(item["def"])
        values[f"fuel_{index}_total"] = money(number(item["total"]))
    values["fuel_total"] = money(trip.fuel_total)

    values["expense_scales"] = money(category_amount(trip, "scale", "scales"))
    values["expense_tolls"] = money(category_amount(trip, "toll", "tolls"))
    values["expense_repairs"] = money(category_amount(trip, "repair", "repairs"))
    values["expense_lumpers"] = money(category_amount(trip, "lumper", "lumpers"))
    other_rows = [item for item in trip.expense_rows if item["category"].lower() not in {"scale", "scales", "toll", "tolls", "repair", "repairs", "lumper", "lumpers"}]
    other_descriptions = []
    for item in other_rows:
        label = item["description"] or item["category"]
        amount = money(number(item["amount"]))
        other_descriptions.append(f"{label} {amount}".strip())
    values["expense_other"] = "; ".join(other_descriptions)[:70]
    values["expense_total"] = money(trip.expenses)

    for index, item in enumerate(trip.advance_rows[:7], 1):
        values[f"advance_{index}_date"] = short_date(item["date"])
        values[f"advance_{index}_amount"] = money(number(item["amount"]))
        values[f"advance_{index}_type"] = item["type"]
    values["advance_total"] = money(trip.advances)

    values["total_miles_pay"] = money(trip.mileage_pay, blank_zero=False)
    values["total_extra_stops_pay"] = money(trip.extra_stop_pay)
    values["total_layover_pay"] = money(trip.layover_pay)
    values["total_money_received"] = money(trip.advances)
    values["total_money_spent"] = money(trip.expenses)
    values["money_left_on_hand"] = money(trip.advances - trip.expenses)
    values["total_pay_to_driver"] = money(trip.settlement_pay, blank_zero=False)
    values["driver_signature"] = ""
    return {name: value for name, value in values.items() if name in field_names}


def count_widgets(reader: PdfReader) -> int:
    count = 0
    for page in reader.pages:
        for annotation_ref in page.get("/Annots", []):
            annotation = annotation_ref.get_object()
            if annotation.get("/Subtype") == "/Widget":
                count += 1
    return count


def sync_canonical_field_values(writer: PdfWriter, values: dict[str, str]) -> None:
    acroform = writer.root_object.get("/AcroForm")
    if not acroform:
        return
    stack = list(acroform.get_object().get("/Fields", []))
    while stack:
        field_ref = stack.pop()
        field = field_ref.get_object()
        name = clean(field.get("/T"))
        if name in values:
            field[NameObject("/V")] = TextStringObject(values[name])
        stack.extend(field.get("/Kids", []))


def export_pdf(template: Path, trip: Trip, export_dir: Path, overwrite: bool) -> tuple[Path, Path]:
    reader = PdfReader(template)
    if len(reader.pages) != 2:
        raise ValueError("The settlement template must contain exactly two pages.")
    fields = reader.get_fields() or {}
    required = {"driver_name", "trip_number", "beginning_odometer", "ending_odometer", "total_miles_pay", "total_pay_to_driver"}
    missing = required - set(fields)
    if missing:
        raise ValueError(f"Template is missing required fields: {', '.join(sorted(missing))}")
    text_fields = {name for name, field in fields.items() if field.get("/FT") == "/Tx"}
    values = build_field_values(trip, text_fields)

    driver = safe_name(trip.driver)
    base = f"{driver}_Trip_{safe_name(trip.trip_number)}_Settlement"
    fillable_path = export_dir / f"{base}_Fillable.pdf"
    flat_path = export_dir / f"{base}.pdf"
    if not overwrite:
        for path in (fillable_path, flat_path):
            if path.exists():
                response = input(f"{path.name} already exists. Overwrite it? [y/N]: ").strip().lower()
                if response not in {"y", "yes"}:
                    raise FileExistsError(f"Export cancelled to protect {path.name}.")

    export_dir.mkdir(parents=True, exist_ok=True)
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    writer.update_page_form_field_values(None, values, auto_regenerate=False)
    # This template contains a canonical field tree plus same-named page
    # widgets. pypdf updates the widgets, so explicitly synchronize the
    # canonical values without reattaching or duplicating fields.
    sync_canonical_field_values(writer, values)
    with fillable_path.open("wb") as stream:
        writer.write(stream)

    check_reader = PdfReader(fillable_path)
    check_fields = check_reader.get_fields() or {}
    for name, expected in values.items():
        if expected and clean(check_fields.get(name, {}).get("/V")) != expected:
            raise RuntimeError(f"Fillable PDF validation failed for {name}.")

    flat_reader = PdfReader(fillable_path)
    flat_writer = PdfWriter()
    flat_writer.clone_document_from_reader(flat_reader)
    current_fields = flat_writer.get_fields() or {}
    current_values = {
        name: field.get("/V", "/Off" if field.get("/FT") == "/Btn" else "")
        for name, field in current_fields.items()
    }
    flat_writer.update_page_form_field_values(None, current_values, auto_regenerate=False, flatten=True)
    flat_writer.remove_annotations(subtypes="/Widget")
    flat_writer.root_object.pop(NameObject("/AcroForm"), None)
    with flat_path.open("wb") as stream:
        flat_writer.write(stream)

    final_reader = PdfReader(flat_path)
    if len(final_reader.pages) != 2:
        raise RuntimeError("Flattened PDF does not contain exactly two pages.")
    if final_reader.get_fields():
        raise RuntimeError("Flattened PDF still contains an AcroForm field tree.")
    if count_widgets(final_reader) != 0:
        raise RuntimeError("Flattened PDF still contains widget annotations.")
    return fillable_path, flat_path


def confirm_trip(trip: Trip, assume_yes: bool) -> None:
    print(f"\nTrip {trip.trip_number}")
    print(f"  Actual odometer miles: {trip.actual_miles:,}")
    print(f"  Paid miles:            {trip.paid_miles:,}")
    print(f"  State miles:           {trip.state_miles:,}")
    print(f"  Mileage pay:           ${trip.mileage_pay:,.2f}")
    print(f"  Gross pay:             ${trip.gross_pay:,.2f}")
    print(f"  Settlement pay:        ${trip.settlement_pay:,.2f}")
    if trip.warnings:
        print("  Warnings:")
        for warning in trip.warnings:
            print(f"    - {warning}")
        if not assume_yes:
            response = input("Continue without changing these values? [y/N]: ").strip().lower()
            if response not in {"y", "yes"}:
                raise RuntimeError("Export cancelled. No mileage values were changed.")


def combine_pdfs(paths: list[Path], output: Path) -> None:
    writer = PdfWriter()
    for path in paths:
        writer.append(str(path))
    with output.open("wb") as stream:
        writer.write(stream)


def main() -> int:
    parser = argparse.ArgumentParser(description="Export United Freight settlement PDFs from the trip-pay workbook.")
    parser.add_argument("--workbook", default="Dustin_Weekly_Trip_Pay_Calculator.xlsx")
    parser.add_argument("--template", default="Settlement_Template.pdf")
    parser.add_argument("--output-dir", default="Exports")
    parser.add_argument("--trip", help="Export this trip number instead of Weekly Pay N4.")
    parser.add_argument("--all", action="store_true", help="Export every completed trip listed on Weekly Pay.")
    parser.add_argument("--combine", action="store_true", help="Create a combined weekly PDF when exporting all trips.")
    parser.add_argument("--yes", action="store_true", help="Accept mileage and receipt warnings.")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing exports without prompting.")
    args = parser.parse_args()

    workbook_path = Path(args.workbook).resolve()
    template_path = Path(args.template).resolve()
    output_dir = Path(args.output_dir).resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    if not template_path.exists():
        raise FileNotFoundError(f"PDF template not found: {template_path}")

    data = WorkbookData(workbook_path)
    trip_numbers = data.all_trip_numbers() if args.all else [args.trip or data.selected_trip()]
    trip_numbers = [trip for trip in trip_numbers if trip]
    if not trip_numbers:
        raise ValueError("No trip was selected and no completed trips were found.")

    flattened: list[Path] = []
    for trip_number in trip_numbers:
        trip = data.trip(trip_number)
        confirm_trip(trip, args.yes)
        fillable, flat = export_pdf(template_path, trip, output_dir, args.overwrite)
        flattened.append(flat)
        print(f"  Fillable: {fillable}")
        print(f"  Ready to email: {flat}")

    if args.all and args.combine and flattened:
        week_start = as_date(data.weekly["E4"].value)
        date_text = week_start.isoformat() if week_start else "week"
        combined = output_dir / f"{safe_name(clean(data.weekly['B4'].value))}_Week_{date_text}_Settlements.pdf"
        combine_pdfs(flattened, combined)
        print(f"  Combined weekly PDF: {combined}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
